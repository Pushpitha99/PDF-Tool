import os
import asyncio
from docx import Document
from services.pdf_service import merge_word_documents


def create_sample_docs():
    """Generates sample Word documents for testing if they do not exist."""
    files = ["sample1.docx", "sample2.docx", "sample3.docx"]
    for idx, filename in enumerate(files, 1):
        if not os.path.exists(filename):
            doc = Document()
            doc.add_heading(f"Document {idx} Title", level=1)
            doc.add_paragraph(f"This is sample paragraph text from document #{idx}.")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = f"Doc {idx} Cell 1"
            table.cell(0, 1).text = f"Doc {idx} Cell 2"
            doc.save(filename)
            print(f"Created sample file: {filename}")
    return files


async def run_test():
    # Ensure test files exist
    files = create_sample_docs()
    output_path = "merged_result.docx"

    print("\n--- Starting Word Document Merge Test ---")
    print(f"Merging files: {files}")
    
    # Merge documents using pdf_service logic
    result = await merge_word_documents(files, output_path)
    
    print("------------------------------------------")
    print(f"SUCCESS: Merged document created at '{result}'")
    print(f"File size: {os.path.getsize(result)} bytes")


if __name__ == "__main__":
    asyncio.run(run_test())
