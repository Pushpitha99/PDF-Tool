import io
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from docx import Document
from reportlab.pdfgen import canvas

from main import app


@pytest_asyncio.fixture
async def client():
    """Async HTTP client fixture configured to query the FastAPI app directly."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


def _generate_dummy_pdf_bytes(title: str = "Test Document") -> bytes:
    """Helper function that generates a valid in-memory PDF document."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(100, 750, title)
    c.drawString(100, 730, "Automated QA Test Document Content.")
    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def _generate_dummy_docx_bytes(title: str = "Test Word Document") -> bytes:
    """Helper function that generates a valid in-memory Word .docx document."""
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Automated QA Test Word Document Content.")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@pytest.fixture
def valid_pdf_1():
    """Fixture providing first valid in-memory PDF tuple."""
    return ("doc1.pdf", _generate_dummy_pdf_bytes("PDF Sample One"), "application/pdf")


@pytest.fixture
def valid_pdf_2():
    """Fixture providing second valid in-memory PDF tuple."""
    return ("doc2.pdf", _generate_dummy_pdf_bytes("PDF Sample Two"), "application/pdf")


@pytest.fixture
def valid_docx_1():
    """Fixture providing first valid in-memory Word .docx tuple."""
    return ("word1.docx", _generate_dummy_docx_bytes("Word Sample One"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@pytest.fixture
def valid_docx_2():
    """Fixture providing second valid in-memory Word .docx tuple."""
    return ("word2.docx", _generate_dummy_docx_bytes("Word Sample Two"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@pytest.fixture
def corrupted_file():
    """Fixture providing a text file disguised as a PDF."""
    return ("fake.pdf", b"THIS IS NOT A VALID PDF CONTENT AT ALL!", "application/pdf")


@pytest.fixture
def empty_file():
    """Fixture providing a 0-byte empty file."""
    return ("empty.pdf", b"", "application/pdf")
